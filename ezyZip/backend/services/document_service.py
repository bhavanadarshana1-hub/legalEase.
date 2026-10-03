from io import BytesIO
from pathlib import Path
from typing import Optional

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# ============================================================
# TXT EXPORT
# ============================================================

def format_txt(text: str) -> bytes:
    """
    Convert document text into a UTF-8 TXT file.
    """

    if text is None:
        text = ""

    return text.encode("utf-8")


# ============================================================
# DOCX EXPORT
# ============================================================

def format_docx(
    text: str,
    document_type: str,
    terms=None,
    brand_name: str = "LegalEase",
    logo_path: Optional[str] = None,
) -> bytes:
    """
    Generate a DOCX document.

    Returns:
        bytes: DOCX file contents
    """

    document = Document()

    # --------------------------------------------------------
    # PAGE MARGINS
    # --------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # --------------------------------------------------------
    # LOGO
    # --------------------------------------------------------

    if logo_path:

        logo = Path(logo_path)

        if logo.exists():

            try:

                paragraph = document.add_paragraph()

                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                run = paragraph.add_run()

                run.add_picture(
                    str(logo),
                    width=Inches(1.2),
                )

            except Exception:
                pass

    # --------------------------------------------------------
    # BRAND
    # --------------------------------------------------------

    brand_paragraph = document.add_paragraph()

    brand_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    brand_run = brand_paragraph.add_run(
        brand_name or "LegalEase"
    )

    brand_run.bold = True
    brand_run.font.size = Pt(18)

    # --------------------------------------------------------
    # DOCUMENT TYPE
    # --------------------------------------------------------

    title_paragraph = document.add_paragraph()

    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title_paragraph.add_run(
        document_type or "Legal Document"
    )

    title_run.bold = True
    title_run.font.size = Pt(15)

    # --------------------------------------------------------
    # SPACING
    # --------------------------------------------------------

    document.add_paragraph()

    # --------------------------------------------------------
    # DOCUMENT CONTENT
    # --------------------------------------------------------

    if not text:
        text = "No document content was provided."

    lines = text.splitlines()

    for line in lines:

        line = line.strip()

        if not line:

            document.add_paragraph()

            continue

        # Markdown headings
        if line.startswith("### "):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line[4:].strip()
            )

            run.bold = True
            run.font.size = Pt(13)

        elif line.startswith("## "):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line[3:].strip()
            )

            run.bold = True
            run.font.size = Pt(14)

        elif line.startswith("# "):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line[2:].strip()
            )

            run.bold = True
            run.font.size = Pt(15)

        elif line.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:].strip()
            )

        elif line.startswith("* "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:].strip()
            )

        else:

            paragraph = document.add_paragraph()

            paragraph.add_run(line)

    # --------------------------------------------------------
    # RETURN DOCX
    # --------------------------------------------------------

    buffer = BytesIO()

    document.save(buffer)

    data = buffer.getvalue()

    buffer.close()

    return data


# ============================================================
# PDF EXPORT
# ============================================================

def format_pdf(
    text: str,
    document_type: str,
    terms=None,
    brand_name: str = "LegalEase",
) -> bytes:
    """
    Generate a PDF document using ReportLab.

    Returns:
        bytes: PDF file contents
    """

    # --------------------------------------------------------
    # SAFETY
    # --------------------------------------------------------

    if text is None:
        text = ""

    if document_type is None:
        document_type = "Legal Document"

    if brand_name is None:
        brand_name = "LegalEase"

    # --------------------------------------------------------
    # PDF BUFFER
    # --------------------------------------------------------

    buffer = BytesIO()

    # --------------------------------------------------------
    # PDF DOCUMENT
    # --------------------------------------------------------

    pdf = SimpleDocTemplate(

        buffer,

        pagesize=A4,

        rightMargin=20 * mm,

        leftMargin=20 * mm,

        topMargin=20 * mm,

        bottomMargin=20 * mm,

        title=str(document_type),

        author=str(brand_name),
    )

    # --------------------------------------------------------
    # STYLES
    # --------------------------------------------------------

    styles = getSampleStyleSheet()

    brand_style = ParagraphStyle(

        "LegalEaseBrand",

        parent=styles["Title"],

        fontName="Helvetica-Bold",

        fontSize=18,

        leading=22,

        alignment=TA_CENTER,

        textColor=colors.HexColor(
            "#0B3B4C"
        ),

        spaceAfter=5 * mm,
    )


    title_style = ParagraphStyle(

        "LegalEaseTitle",

        parent=styles["Heading1"],

        fontName="Helvetica-Bold",

        fontSize=15,

        leading=19,

        alignment=TA_CENTER,

        textColor=colors.HexColor(
            "#222222"
        ),

        spaceAfter=8 * mm,
    )


    heading_style = ParagraphStyle(

        "LegalEaseHeading",

        parent=styles["Heading2"],

        fontName="Helvetica-Bold",

        fontSize=12,

        leading=16,

        textColor=colors.HexColor(
            "#0B3B4C"
        ),

        spaceBefore=5 * mm,

        spaceAfter=3 * mm,
    )


    body_style = ParagraphStyle(

        "LegalEaseBody",

        parent=styles["BodyText"],

        fontName="Helvetica",

        fontSize=10.5,

        leading=15,

        textColor=colors.HexColor(
            "#222222"
        ),

        spaceAfter=4 * mm,
    )


    bullet_style = ParagraphStyle(

        "LegalEaseBullet",

        parent=body_style,

        leftIndent=8 * mm,

        firstLineIndent=-4 * mm,

        bulletIndent=3 * mm,
    )

    # --------------------------------------------------------
    # STORY
    # --------------------------------------------------------

    story = []

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    story.append(
        Paragraph(
            str(brand_name),
            brand_style,
        )
    )


    story.append(
        Paragraph(
            str(document_type),
            title_style,
        )
    )


    # --------------------------------------------------------
    # SEPARATOR
    # --------------------------------------------------------

    separator = Table(
        [[""]],
        colWidths=[170 * mm],
        rowHeights=[0.5 * mm],
    )

    separator.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor("#36C2B4"),
                ),
            ]
        )
    )

    story.append(separator)

    story.append(
        Spacer(
            1,
            8 * mm
        )
    )

    # --------------------------------------------------------
    # DOCUMENT CONTENT
    # --------------------------------------------------------

    lines = text.splitlines()

    for line in lines:

        line = line.strip()

        # Empty line
        if not line:

            story.append(
                Spacer(
                    1,
                    3 * mm
                )
            )

            continue

        # ----------------------------------------------------
        # HEADINGS
        # ----------------------------------------------------

        if line.startswith("### "):

            heading = line[4:].strip()

            story.append(
                Paragraph(
                    _escape_pdf_text(
                        heading
                    ),
                    heading_style,
                )
            )

            continue


        if line.startswith("## "):

            heading = line[3:].strip()

            story.append(
                Paragraph(
                    _escape_pdf_text(
                        heading
                    ),
                    heading_style,
                )
            )

            continue


        if line.startswith("# "):

            heading = line[2:].strip()

            story.append(
                Paragraph(
                    _escape_pdf_text(
                        heading
                    ),
                    heading_style,
                )
            )

            continue

        # ----------------------------------------------------
        # BULLETS
        # ----------------------------------------------------

        if line.startswith("- "):

            bullet_text = line[2:].strip()

            story.append(
                Paragraph(
                    "• "
                    + _escape_pdf_text(
                        bullet_text
                    ),
                    bullet_style,
                )
            )

            continue


        if line.startswith("* "):

            bullet_text = line[2:].strip()

            story.append(
                Paragraph(
                    "• "
                    + _escape_pdf_text(
                        bullet_text
                    ),
                    bullet_style,
                )
            )

            continue

        # ----------------------------------------------------
        # NORMAL TEXT
        # ----------------------------------------------------

        safe_text = _escape_pdf_text(
            line
        )

        # Basic Markdown bold
        safe_text = _convert_basic_markdown(
            safe_text
        )

        story.append(
            Paragraph(
                safe_text,
                body_style,
            )
        )

    # --------------------------------------------------------
    # EMPTY DOCUMENT
    # --------------------------------------------------------

    if not lines:

        story.append(
            Paragraph(
                "No document content was provided.",
                body_style,
            )
        )

    # --------------------------------------------------------
    # BUILD PDF
    # --------------------------------------------------------

    pdf.build(story)

    # --------------------------------------------------------
    # GET BYTES
    # --------------------------------------------------------

    pdf_bytes = buffer.getvalue()

    buffer.close()

    return pdf_bytes


# ============================================================
# PDF TEXT ESCAPING
# ============================================================

def _escape_pdf_text(text: str) -> str:
    """
    Escape characters that ReportLab's Paragraph parser
    treats as markup.
    """

    if text is None:
        return ""

    text = str(text)

    text = (
        text
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    return text


# ============================================================
# BASIC MARKDOWN CONVERSION
# ============================================================

def _convert_basic_markdown(
    text: str
) -> str:
    """
    Convert simple **bold** Markdown into
    ReportLab <b> markup.
    """

    if not text:
        return text

    result = ""

    parts = text.split("**")

    for index, part in enumerate(parts):

        if index % 2 == 1:

            result += (
                "<b>"
                + part
                + "</b>"
            )

        else:

            result += part

    return result